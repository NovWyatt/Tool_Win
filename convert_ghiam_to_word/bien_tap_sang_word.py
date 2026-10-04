# -*- coding: utf-8 -*-
"""
Dựng file Word hoàn chỉnh từ bản đã biên tập (văn bản thuần có vài ký hiệu đơn giản).

Dùng: .venv\\Scripts\\python.exe bien_tap_sang_word.py <bản biên tập .txt> <file .docx ra>

Ký hiệu trong bản biên tập (mỗi dòng một đoạn):
  % Dòng tiêu đề                 in đậm, căn giữa (có thể nhiều dòng)
  %% Dòng phụ đề                 in nghiêng, căn giữa
  : Nhãn: giá trị                dòng thông tin, nhãn in đậm
  > Ghi chú                      chữ nghiêng nhỏ màu xám; sau ghi chú cuối có đường kẻ ngang
  [[MUC_LUC]]                    mục lục (Word cập nhật khi mở, hoặc cập nhật bằng Word COM)
  ## Tiêu đề mục / ### Tiêu đề mục con
  [00:12:34] **Người nói:** lời nói...     đoạn văn, mốc thời gian và nhãn người nói đều không bắt buộc
Trong đoạn văn: ==chỗ cần đối chiếu== tô vàng, *ghi chú của người biên tập* in nghiêng,
**chữ đậm**, [không nghe rõ] in nghiêng màu xám.
"""

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from chuyen_doi import KHONG_RO, XAM, dat_font, ke_duong_duoi, tao_tai_lieu, them_truong  # noqa: E402

MOC_GIO = re.compile(r"^\[(\d{1,2}:\d{2}:\d{2})\]\s*")
KY_HIEU = re.compile(r"(==.+?==|\*\*.+?\*\*|\*[^*\s][^*]*?\*|" + re.escape(KHONG_RO) + ")")


def them_noi_dung(p, van_ban):
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import RGBColor

    for phan in KY_HIEU.split(van_ban):
        if not phan:
            continue
        if phan.startswith("==") and phan.endswith("==") and len(phan) > 4:
            p.add_run(phan[2:-2]).font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif phan.startswith("**") and phan.endswith("**") and len(phan) > 4:
            p.add_run(phan[2:-2]).bold = True
        elif phan == KHONG_RO:
            r = p.add_run(phan)
            r.italic = True
            r.font.color.rgb = RGBColor(*XAM)
        elif phan.startswith("*") and phan.endswith("*") and len(phan) > 2:
            p.add_run(phan[1:-1]).italic = True
        else:
            p.add_run(phan)


def them_muc_luc(doc):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("MỤC LỤC")
    r.bold = True
    them_truong(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', Pt(13))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def dung_word(nguon, dich):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = tao_tai_lieu()
    for ten, co_chu in (("Heading 1", Pt(13)), ("Heading 2", Pt(13))):
        kieu = doc.styles[ten]
        dat_font(kieu, co_chu, dam=True)
        kieu.font.italic = ten == "Heading 2"
        kieu.font.color.rgb = RGBColor(0, 0, 0)
        kieu.paragraph_format.space_before = Pt(12)
        kieu.paragraph_format.space_after = Pt(6)
        kieu.paragraph_format.keep_with_next = True
    # Dòng mục lục gọn hơn đoạn văn thường (Word dùng hai kiểu có sẵn tên "toc 1", "toc 2" khi cập nhật mục lục).
    for ten, thut_le in (("toc 1", 0), ("toc 2", 0.5)):
        kieu = doc.styles.add_style(ten, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
        kieu.style_id = ten.replace(" ", "").upper()  # "TOC1": đúng mã kiểu có sẵn, Word mới nhận là kiểu mục lục
        kieu.base_style = doc.styles["Normal"]
        kieu.paragraph_format.space_after = Pt(2)
        kieu.paragraph_format.line_spacing = 1.0
        kieu.paragraph_format.left_indent = Cm(thut_le)

    xam = RGBColor(*XAM)
    ghi_chu_cuoi = None
    for dong in Path(nguon).read_text(encoding="utf-8-sig").splitlines():
        dong = dong.strip()
        if not dong:
            continue
        if not dong.startswith(">") and ghi_chu_cuoi is not None:
            ke_duong_duoi(ghi_chu_cuoi)
            ghi_chu_cuoi = None

        if dong.startswith("%% "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            p.add_run(dong[3:]).italic = True
        elif dong.startswith("% "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(dong[2:])
            r.bold = True
            r.font.size = Pt(14)
        elif dong.startswith(": "):
            nhan, _, gia_tri = dong[2:].partition(":")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.add_run(nhan + ":").bold = True
            them_noi_dung(p, gia_tri)
        elif dong.startswith(">"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(dong[1:].strip())
            r.italic = True
            r.font.size = Pt(11)
            r.font.color.rgb = xam
            ghi_chu_cuoi = p
        elif dong == "[[MUC_LUC]]":
            them_muc_luc(doc)
        elif dong.startswith("### "):
            doc.add_paragraph(dong[4:], style="Heading 2")
        elif dong.startswith("## "):
            doc.add_paragraph(dong[3:], style="Heading 1")
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            moc = MOC_GIO.match(dong)
            if moc:
                r = p.add_run(f"[{moc[1]}] ")
                r.font.size = Pt(10)
                r.font.color.rgb = xam
                dong = dong[moc.end():]
            them_noi_dung(p, dong)
    if ghi_chu_cuoi is not None:
        ke_duong_duoi(ghi_chu_cuoi)
    doc.save(str(dich))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    dung_word(sys.argv[1], sys.argv[2])
    print(f"Đã tạo {sys.argv[2]}")
