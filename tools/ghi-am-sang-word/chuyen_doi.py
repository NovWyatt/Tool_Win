# -*- coding: utf-8 -*-
"""
CHUYỂN GHI ÂM THÀNH VĂN BẢN (tiếng Việt) — chạy offline trên máy bằng Whisper large-v3.

Cách dùng:
  - Bấm đúp CHUYEN_DOI_GHI_AM.bat: xử lý mọi file ghi âm trong thư mục này chưa có kết quả.
  - Kéo thả file ghi âm vào CHUYEN_DOI_GHI_AM.bat: xử lý đúng các file đó.
  - Dòng lệnh:
      .venv\\Scripts\\python.exe chuyen_doi.py [file ...] [--goi-y "tên riêng, thuật ngữ"] [--cpu] [--lam-lai] [--ra thư_mục]

Kết quả nằm trong thư mục ket_qua\\ (hoặc thư mục chọn bằng --ra), đặt tên theo tên file ghi âm kèm cả đuôi
để "hop.m4a" và "hop.mp4" không ghi đè nhau:
  hop.m4a - ban tho.docx    văn bản Word, chia đoạn theo chỗ ngắt nghỉ, có mốc thời gian
  hop.m4a - ban tho.txt     từng câu kèm mốc thời gian bắt đầu → kết thúc
  _du_lieu\\hop.m4a.json     dữ liệu chi tiết (độ tin cậy từng câu), dùng khi biên tập lại
Kết quả bản trước đặt tên bỏ đuôi ("hop - ban tho.docx") vẫn được nhận ra nên không bị làm lại.

Mã thoát: 0 = xong hết, 1 = có file bị lỗi, 2 = không có file nào để xử lý.
"""

import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

THU_MUC = Path(__file__).resolve().parent
MO_HINH = THU_MUC / "models" / "faster-whisper-large-v3"
KET_QUA = THU_MUC / "ket_qua"           # đổi bằng --ra (xem dat_thu_muc_ket_qua)
DU_LIEU = KET_QUA / "_du_lieu"
DUOI_KET_QUA = " - ban tho"

TAN_SO = 16000
DUOI_AM_THANH = {
    ".mp3", ".m4a", ".wav", ".aac", ".ogg", ".oga", ".opus", ".flac", ".wma", ".amr",
    ".3gp", ".3ga", ".m4b", ".caf", ".aif", ".aiff", ".weba", ".mka",
    ".mp4", ".mov", ".mkv", ".webm", ".avi", ".mpg", ".mpeg", ".ts",
}

# Câu "ảo giác" quen thuộc của Whisper tiếng Việt: gặp tiếng ồn, tiếng xì xào hay nhạc nền là
# nó bịa ra lời chào cuối video YouTube thay cho nội dung thật.
AO_GIAC = re.compile(
    r"ghiền mì gõ|la la school|subscribe|đăng ký kênh|ủng hộ kênh|theo dõi kênh|không bỏ lỡ những video|"
    r"video hấp dẫn|video tiếp theo|bấm chuông|like và share|cảm ơn các bạn đã (theo dõi|xem)|"
    r"đã theo dõi và hẹn gặp lại",
    re.IGNORECASE,
)
CHI_CO_CHAO = re.compile(r"^\W*(xin )?(cảm ơn|hẹn gặp lại)[^.!?]{0,25}\W*$", re.IGNORECASE)

SO_CHUM = 5                # beam size: số phương án câu mô hình cân nhắc song song
VAD_NGUONG = 0.2           # ngưỡng dò tiếng nói; ghi âm qua loa hội trường cần ngưỡng thấp
CUA_SO_VAD = 120.0         # dò tiếng nói theo từng cửa sổ chừng này giây (xem tim_khuc_tieng_noi)
KHUC_TOI_DA = 25.0         # mỗi khúc gửi cho mô hình dài tối đa chừng này giây
KHOANG_LANG_GOP = 2.0      # hai đoạn nói cách nhau không quá chừng này giây thì gộp chung khúc
MAT_DO_TOI_THIEU = 2.5     # số ký tự mỗi giây nói; thấp hơn là mô hình đã bỏ sót lời
NGUONG_KEM_CHAC = -0.8     # avg_logprob thấp hơn mức này thì tô vàng để nghe lại
NGAT_DOAN_GIAY = 2.0       # im lặng từ chừng này giây trở lên thì sang đoạn mới
DOAN_TOI_DA_KY_TU = 900    # đoạn dài quá mức này thì ngắt ở cuối câu gần nhất
KHONG_RO = "[không nghe rõ]"


def nap_thu_vien_cuda():
    """Cho Windows tìm thấy cuBLAS/cuDNN cài bằng pip (.venv\\Lib\\site-packages\\nvidia\\*\\bin)."""
    if sys.platform != "win32":
        return
    goc = Path(sys.prefix) / "Lib" / "site-packages" / "nvidia"
    for thu_muc_bin in goc.glob("*/bin"):
        os.add_dll_directory(str(thu_muc_bin))
        os.environ["PATH"] = str(thu_muc_bin) + os.pathsep + os.environ.get("PATH", "")


def gio(giay):
    giay = max(0, int(giay))
    return f"{giay // 3600:02d}:{giay % 3600 // 60:02d}:{giay % 60:02d}"


NHAN_MO_HINH = ""          # thêm vào tên file kết quả khi dùng mô hình khác mặc định


def dat_thu_muc_ket_qua(thu_muc):
    global KET_QUA, DU_LIEU
    KET_QUA = Path(thu_muc).resolve()
    DU_LIEU = KET_QUA / "_du_lieu"


def file_ket_qua(file_am_thanh, duoi):
    # Giữ cả đuôi file gốc: "hop.m4a" và "hop.mp4" cùng thư mục thì kết quả không ghi đè nhau.
    return KET_QUA / f"{file_am_thanh.name}{DUOI_KET_QUA}{NHAN_MO_HINH}{duoi}"


def file_du_lieu(file_am_thanh):
    return DU_LIEU / f"{file_am_thanh.name}{NHAN_MO_HINH}.json"


def da_co_ket_qua(file_am_thanh):
    if file_ket_qua(file_am_thanh, ".docx").exists():
        return True
    # Bản 2026.10 đặt tên bỏ đuôi ("hop - ban tho.docx"). Chỉ coi là kết quả của file này khi dữ liệu kèm theo ghi
    # đúng tên file ghi âm, để "hop.mp4" không bị bỏ qua chỉ vì "hop.m4a" đã làm xong.
    if not (KET_QUA / f"{file_am_thanh.stem}{DUOI_KET_QUA}{NHAN_MO_HINH}.docx").exists():
        return False
    try:
        du_lieu = json.loads((DU_LIEU / f"{file_am_thanh.stem}{NHAN_MO_HINH}.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return False
    return isinstance(du_lieu, dict) and du_lieu.get("tep") == file_am_thanh.name


def tim_file_dau_vao(ds_duong_dan, lam_lai):
    if ds_duong_dan:
        ung_vien = []
        for p in map(Path, ds_duong_dan):
            if p.is_dir():
                ung_vien += sorted(x for x in p.iterdir() if x.suffix.lower() in DUOI_AM_THANH)
            elif p.is_file():
                ung_vien.append(p)
            else:
                print(f"Không tìm thấy: {p}")
    else:
        ung_vien = sorted(x for x in THU_MUC.iterdir() if x.is_file() and x.suffix.lower() in DUOI_AM_THANH)

    ds = []
    for f in ung_vien:
        if not lam_lai and da_co_ket_qua(f):
            print(f"Bỏ qua (đã có kết quả): {f.name}   — muốn làm lại thì thêm --lam-lai")
        else:
            ds.append(f.resolve())
    return ds


def doc_am_thanh(duong_dan):
    """Giải mã file âm thanh/video thành mảng mono 16 kHz."""
    import numpy as np
    from faster_whisper.audio import decode_audio

    try:
        with open(duong_dan, "rb") as f:
            return decode_audio(f, sampling_rate=TAN_SO)
    except Exception as loi_giai_ma:
        ffmpeg = shutil.which("ffmpeg")
        if not ffmpeg:
            raise RuntimeError(f"không đọc được file âm thanh ({loi_giai_ma})") from loi_giai_ma
        kq = subprocess.run(
            [ffmpeg, "-nostdin", "-v", "error", "-i", str(duong_dan),
             "-f", "s16le", "-ac", "1", "-ar", str(TAN_SO), "-"],
            capture_output=True,
        )
        if kq.returncode != 0 or not kq.stdout:
            chi_tiet = kq.stderr.decode("utf-8", "replace").strip()[-400:]
            raise RuntimeError(f"không đọc được file âm thanh: {chi_tiet}") from loi_giai_ma
        return np.frombuffer(kq.stdout, np.int16).astype(np.float32) / 32768.0


def nap_mo_hinh(dung_cpu, thu_muc_mo_hinh=None):
    from faster_whisper import WhisperModel

    duong_dan = str(thu_muc_mo_hinh or MO_HINH)
    if not dung_cpu:
        try:
            return WhisperModel(duong_dan, device="cuda", compute_type="int8_float16"), "card đồ họa NVIDIA"
        except Exception as loi:
            print(f"Không dùng được card đồ họa ({loi}).\nChuyển sang CPU — sẽ chậm hơn nhiều.")
    so_luong = max(1, min(8, os.cpu_count() or 4))
    return WhisperModel(duong_dan, device="cpu", compute_type="int8", cpu_threads=so_luong), "CPU"


def tim_khuc_tieng_noi(am_thanh):
    """Dò chỗ có tiếng nói rồi gộp thành các khúc ≤ KHUC_TOI_DA giây, cắt ở chỗ ngắt nghỉ.

    Bộ dò (Silero) mang trạng thái suốt file: với ghi âm dài, ồn liên tục nó bị "trôi" và bỏ qua cả
    chục phút tiếng nói. Vì vậy dò từng cửa sổ CUA_SO_VAD giây, mỗi cửa sổ bắt đầu lại từ đầu.
    """
    from faster_whisper.vad import VadOptions, get_speech_timestamps

    tuy_chon = VadOptions(threshold=VAD_NGUONG, min_silence_duration_ms=500, speech_pad_ms=300,
                          max_speech_duration_s=KHUC_TOI_DA)
    buoc = int(CUA_SO_VAD * TAN_SO)
    doan_noi = []
    for dau in range(0, len(am_thanh), buoc):
        for t in get_speech_timestamps(am_thanh[dau:dau + buoc], tuy_chon):
            doan_noi.append(((dau + t["start"]) / TAN_SO, (dau + t["end"]) / TAN_SO))

    khuc = []
    for bd, kt in doan_noi:
        if khuc and bd - khuc[-1]["kt"] <= KHOANG_LANG_GOP and kt - khuc[-1]["bd"] <= KHUC_TOI_DA:
            khuc[-1]["kt"] = kt
            khuc[-1]["noi"] += kt - bd
        else:
            khuc.append({"bd": bd, "kt": kt, "noi": kt - bd})
    return khuc


def diem_cat_giua(am_thanh, bd, kt):
    """Chỗ nhỏ tiếng nhất quanh giữa khúc — chia đôi ở đó để khỏi cắt ngang lời."""
    import numpy as np

    a, b = int((bd + (kt - bd) * 0.3) * TAN_SO), int((bd + (kt - bd) * 0.7) * TAN_SO)
    khung = TAN_SO // 10
    so_khung = (b - a) // khung
    if so_khung < 2:
        return (bd + kt) / 2
    nang_luong = (am_thanh[a:a + so_khung * khung].reshape(so_khung, khung) ** 2).mean(axis=1)
    return (a + int(np.argmin(nang_luong)) * khung + khung // 2) / TAN_SO


def chuan_hoa(s):
    return re.sub(r"\W+", " ", s.lower()).strip()


def chep_lai(cau, loi_nhac):
    """Câu có phải chép gần nguyên từ lời nhắc (ngữ cảnh khúc trước) không."""
    from difflib import SequenceMatcher

    if len(cau) < 15 or not loi_nhac:
        return False
    khop = SequenceMatcher(None, cau, loi_nhac, autojunk=False).find_longest_match(0, len(cau), 0, len(loi_nhac))
    return khop.size / len(cau) >= 0.6


def danh_gia(doan_noi, giay_noi, loi_nhac):
    """Bỏ câu bịa và câu chép lại lời nhắc; cho biết kết quả có dùng được không."""
    nhac = chuan_hoa(loi_nhac or "")
    giu = []
    for s in doan_noi:
        ct = chuan_hoa(s.text)
        if not ct or AO_GIAC.search(s.text) or chep_lai(ct, nhac):
            continue
        giu.append(s)
    if len(giu) == 1 and giay_noi >= 6 and CHI_CO_CHAO.match(giu[0].text.strip()):
        giu = []  # cả khúc dài mà chỉ ra một câu "cảm ơn…" thì gần như chắc là bịa
    so_chu = sum(len(s.text.strip()) for s in giu)
    lap_lai = any(s.compression_ratio > 2.4 for s in giu)
    dat = bool(giu) and not lap_lai and (giay_noi < 6 or so_chu / giay_noi >= MAT_DO_TOI_THIEU)
    # Ước lượng số ký tự đúng: dùng để chọn phương án đỡ nhất khi không lần thử nào đạt.
    diem = sum(len(s.text.strip()) * 2.718 ** s.avg_logprob for s in giu)
    return giu, dat, diem


def sang_cau(doan_noi, bd, kt, kem=False):
    return [{
        "bat_dau": round(bd + s.start, 2),
        "ket_thuc": round(min(bd + s.end, kt), 2),
        "van_ban": s.text.strip(),
        "avg_logprob": round(s.avg_logprob, 3),
        "no_speech_prob": round(s.no_speech_prob, 3),
        "compression_ratio": round(s.compression_ratio, 2),
        **({"kem": True} if kem else {}),
    } for s in doan_noi]


def nhan_dang_khuc(model, am_thanh, bd, kt, giay_noi, ngu_canh, goi_y, do_sau=0):
    """Nhận dạng một khúc, tự thử lại khi mô hình bịa câu hoặc bỏ sót. Trả về (các câu, trạng thái)."""
    doan_am = am_thanh[int(bd * TAN_SO):int(kt * TAN_SO)]
    mac_dinh = (0.0, 0.2, 0.4, 0.6, 0.8, 1.0)
    # Lần đầu không kèm ngữ cảnh (kèm thì mô hình hay chép lại câu trước); ngữ cảnh chỉ dùng
    # để kéo mô hình khỏi câu bịa kiểu YouTube khi lần đầu hỏng.
    cac_lan_thu = [(goi_y or None, mac_dinh)]
    if ngu_canh:
        cac_lan_thu.append((f"{goi_y} {ngu_canh}".strip(), mac_dinh))
    cac_lan_thu.append((goi_y or None, (0.4, 0.6, 0.8)))

    do_nhat = None
    for loi_nhac, nhiet_do in cac_lan_thu:
        doan_noi, _ = model.transcribe(
            doan_am, language="vi", beam_size=SO_CHUM, vad_filter=False,
            condition_on_previous_text=False, initial_prompt=loi_nhac, temperature=nhiet_do,
        )
        giu, dat, diem = danh_gia(list(doan_noi), giay_noi, loi_nhac)
        if dat:
            return sang_cau(giu, bd, kt), "tot"
        if giu and (do_nhat is None or diem > do_nhat[1]):
            do_nhat = (giu, diem)

    if kt - bd >= 8 and do_sau < 2:
        giua = diem_cat_giua(am_thanh, bd, kt)
        ti_le = (giua - bd) / (kt - bd)
        trai, _ = nhan_dang_khuc(model, am_thanh, bd, giua, giay_noi * ti_le, ngu_canh, goi_y, do_sau + 1)
        ngu_canh_moi = cap_nhat_ngu_canh(ngu_canh, trai)
        phai, _ = nhan_dang_khuc(model, am_thanh, giua, kt, giay_noi * (1 - ti_le), ngu_canh_moi, goi_y, do_sau + 1)
        return trai + phai, "chia_doi"
    if do_nhat:
        return sang_cau(do_nhat[0], bd, kt, kem=True), "kem"
    if giay_noi >= 3:
        return [{"bat_dau": round(bd, 2), "ket_thuc": round(kt, 2), "van_ban": KHONG_RO,
                 "avg_logprob": -9.0, "khong_ro": True}], "khong_ro"
    return [], "trong"


def cap_nhat_ngu_canh(ngu_canh, cau):
    """Giữ khoảng 200 ký tự cuối đã nhận dạng để làm lời nhắc cho khúc sau."""
    moi = " ".join(c["van_ban"] for c in cau if not c.get("khong_ro"))
    if not moi:
        return ngu_canh
    ghep = f"{ngu_canh} {moi}".strip()
    return ghep if len(ghep) <= 200 else ghep[-200:].split(" ", 1)[-1]


def chuyen_thanh_chu(model, am_thanh, goi_y):
    thoi_luong = len(am_thanh) / TAN_SO
    print("Đang dò các đoạn có tiếng nói...", flush=True)
    cac_khuc = tim_khuc_tieng_noi(am_thanh)
    print(f"Có khoảng {gio(sum(k['noi'] for k in cac_khuc))} tiếng nói, chia thành {len(cac_khuc)} khúc. "
          "Bắt đầu nhận dạng:\n", flush=True)

    ds, ngu_canh = [], ""
    bat_dau = time.monotonic()
    for khuc in cac_khuc:
        cau, trang_thai = nhan_dang_khuc(model, am_thanh, khuc["bd"], khuc["kt"], khuc["noi"], ngu_canh, goi_y)
        khuc["trang_thai"] = trang_thai
        ds += cau
        ngu_canh = cap_nhat_ngu_canh(ngu_canh, cau)

        phan_tram = min(100, int(khuc["kt"] / thoi_luong * 100)) if thoi_luong else 100
        da_chay = time.monotonic() - bat_dau
        if khuc["kt"] >= 120 and da_chay >= 20:
            con_lai = (thoi_luong - khuc["kt"]) / (khuc["kt"] / da_chay)
            uoc_tinh = f"còn ~{max(1, round(con_lai / 60))} phút"
        else:
            uoc_tinh = "đang ước tính"
        for c in cau:
            print(f"[{phan_tram:3d}% | {uoc_tinh}] {gio(c['bat_dau'])}  {c['van_ban']}", flush=True)
    return ds, thoi_luong, cac_khuc


def chia_doan(ds):
    """Gộp các câu liền mạch thành đoạn; ngắt đoạn ở chỗ im lặng lâu hoặc khi đoạn quá dài."""
    ket_qua, doan, truoc = [], None, None
    for m in ds:
        if m.get("loai_bo"):
            continue
        sang_doan_moi = (
            doan is None
            or m["bat_dau"] - truoc["ket_thuc"] >= NGAT_DOAN_GIAY
            or (sum(len(x["van_ban"]) for x in doan) >= DOAN_TOI_DA_KY_TU
                and truoc["van_ban"][-1:] in ".?!…")
        )
        if sang_doan_moi:
            doan = []
            ket_qua.append(doan)
        doan.append(m)
        truoc = m
    return ket_qua


def ghi_txt(duong_dan, ten_file, thoi_luong, ds):
    dong = [
        "NỘI DUNG GHI ÂM — BẢN THÔ DO MÁY TỰ ĐỘNG NHẬN DẠNG",
        f"Tệp ghi âm: {ten_file}",
        f"Thời lượng: {gio(thoi_luong)}",
        f"Ngày chuyển đổi: {dt.datetime.now():%d/%m/%Y %H:%M}",
        "",
    ]
    dong += [f"[{gio(m['bat_dau'])} → {gio(m['ket_thuc'])}] {m['van_ban']}" for m in ds if not m.get("loai_bo")]
    duong_dan.write_text("\n".join(dong) + "\n", encoding="utf-8-sig")


def ke_duong_duoi(doan_van):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    vien = OxmlElement("w:pBdr")
    duoi = OxmlElement("w:bottom")
    for khoa, gia_tri in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", "A0A0A0")):
        duoi.set(qn(khoa), gia_tri)
    vien.append(duoi)
    doan_van._p.get_or_add_pPr().insert_element_before(
        vien, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
        "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def them_truong(doan_van, lenh, co_chu):
    """Chèn trường tự cập nhật của Word (PAGE, NUMPAGES...)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for loai in ("begin", "instr", "separate", "text", "end"):
        r = doan_van.add_run()
        r.font.size = co_chu
        if loai == "text":
            r.text = "1"
            continue
        if loai == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {lenh} "
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), loai)
        r._r.append(el)


XAM = (0x70, 0x70, 0x70)


def dat_font(kieu, co_chu, dam=None):
    """Đặt Times New Roman cho mọi bảng chữ, bỏ font theo theme (theme font được ưu tiên hơn font đặt tay)."""
    from docx.oxml.ns import qn

    kieu.font.name = "Times New Roman"
    kieu.font.size = co_chu
    if dam is not None:
        kieu.font.bold = dam
    rfonts = kieu.element.get_or_add_rPr().rFonts
    for thuoc_tinh in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{thuoc_tinh}"), None)
    for thuoc_tinh in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{thuoc_tinh}"), "Times New Roman")


def tao_tai_lieu():
    """Tài liệu Word khổ A4 theo lề văn bản hành chính, chữ Times New Roman 13, số trang ở chân trang."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)

    kieu = doc.styles["Normal"]
    dat_font(kieu, Pt(13))
    ngon_ngu = OxmlElement("w:lang")
    ngon_ngu.set(qn("w:val"), "vi-VN")  # để Word soát chính tả tiếng Việt, không gạch đỏ cả trang
    kieu.element.get_or_add_rPr().append(ngon_ngu)
    kieu.paragraph_format.space_after = Pt(6)
    kieu.paragraph_format.line_spacing = 1.15

    chan_trang = sec.footer.paragraphs[0]
    chan_trang.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chan_trang.add_run("Trang ").font.size = Pt(10)
    them_truong(chan_trang, "PAGE", Pt(10))
    chan_trang.add_run(" / ").font.size = Pt(10)
    them_truong(chan_trang, "NUMPAGES", Pt(10))
    return doc


def ghi_docx(duong_dan, ten_file, thoi_luong, ds):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    xam = RGBColor(*XAM)
    doc = tao_tai_lieu()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("NỘI DUNG GHI ÂM")
    r.bold = True
    r.font.size = Pt(15)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.add_run("(Bản thô do máy tự động nhận dạng)").italic = True

    for nhan, gia_tri in (
        ("Tệp ghi âm: ", ten_file),
        ("Thời lượng: ", gio(thoi_luong)),
        ("Ngày chuyển đổi: ", f"{dt.date.today():%d/%m/%Y}"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(nhan).bold = True
        p.add_run(gia_tri)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(
        "Lưu ý: máy có thể nghe sai tên riêng, con số và dấu câu. Chỗ tô vàng là chỗ máy kém chắc chắn, "
        "nên nghe lại để đối chiếu. Số trong ngoặc vuông là thời điểm (giờ:phút:giây) trong file ghi âm."
    )
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = xam
    ke_duong_duoi(p)

    for doan in chia_doan(ds):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(f"[{gio(doan[0]['bat_dau'])}] ")
        r.font.size = Pt(10)
        r.font.color.rgb = xam
        for j, m in enumerate(doan):
            if j:
                p.add_run(" ")
            r = p.add_run(m["van_ban"])
            if m.get("khong_ro"):
                r.italic = True
                r.font.color.rgb = xam
            elif m.get("kem") or m["avg_logprob"] < NGUONG_KEM_CHAC:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.save(str(duong_dan))


def xu_ly_mot_file(f, model, thiet_bi, goi_y, thu_muc_mo_hinh):
    bat_dau = time.monotonic()
    am_thanh = doc_am_thanh(f)
    try:
        ds, thoi_luong, cac_khuc = chuyen_thanh_chu(model, am_thanh, goi_y)
    except RuntimeError as loi:
        if thiet_bi == "CPU":
            raise
        print(f"\nCard đồ họa gặp lỗi ({loi}).\nLàm lại file này bằng CPU — sẽ chậm hơn nhiều.\n")
        model, thiet_bi = nap_mo_hinh(True, thu_muc_mo_hinh)
        ds, thoi_luong, cac_khuc = chuyen_thanh_chu(model, am_thanh, goi_y)

    DU_LIEU.mkdir(parents=True, exist_ok=True)
    file_du_lieu(f).write_text(json.dumps({
        "tep": f.name,
        "thoi_luong_giay": round(thoi_luong, 2),
        "mo_hinh": Path(thu_muc_mo_hinh).name,
        "thiet_bi": thiet_bi,
        "goi_y": goi_y,
        "ngay_chuyen": dt.datetime.now().isoformat(timespec="seconds"),
        "khuc": [{k: (round(v, 2) if isinstance(v, float) else v) for k, v in kh.items()} for kh in cac_khuc],
        "cau": ds,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    ghi_txt(file_ket_qua(f, ".txt"), f.name, thoi_luong, ds)
    docx = file_ket_qua(f, ".docx")
    try:
        ghi_docx(docx, f.name, thoi_luong, ds)
    except PermissionError:
        raise RuntimeError(f"không ghi được {docx.name} — file đang mở trong Word? Đóng lại rồi chạy lại.")

    dem = {}
    for kh in cac_khuc:
        dem[kh["trang_thai"]] = dem.get(kh["trang_thai"], 0) + 1
    phut = (time.monotonic() - bat_dau) / 60
    print(f"\nXong trong {phut:.1f} phút. {len(cac_khuc)} khúc: {dem.get('tot', 0)} nghe tốt, "
          f"{dem.get('chia_doi', 0)} phải chia nhỏ, {dem.get('kem', 0)} kém chắc chắn (tô vàng), "
          f"{dem.get('khong_ro', 0)} không nghe rõ.")
    print(f"Kết quả: {docx}")
    return model, thiet_bi


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

    ap = argparse.ArgumentParser(description="Chuyển file ghi âm tiếng Việt thành văn bản Word.")
    ap.add_argument("files", nargs="*", help="file ghi âm (bỏ trống: mọi file ghi âm mới trong thư mục công cụ)")
    ap.add_argument("--goi-y", default="", help='tên riêng/thuật ngữ có trong ghi âm, vd: "Nguyễn Văn A, hợp đồng ủy quyền"')
    ap.add_argument("--cpu", action="store_true", help="chạy bằng CPU thay vì card đồ họa")
    ap.add_argument("--lam-lai", action="store_true", help="làm lại cả file đã có kết quả")
    ap.add_argument("--mo-hinh", default=MO_HINH.name, help="tên thư mục mô hình trong models\\")
    ap.add_argument("--ra", metavar="THU_MUC", help="thư mục kết quả (mặc định: ket_qua trong thư mục công cụ)")
    args = ap.parse_args()

    global NHAN_MO_HINH
    if args.mo_hinh != MO_HINH.name:
        NHAN_MO_HINH = f" ({args.mo_hinh})"
    if args.ra:
        dat_thu_muc_ket_qua(args.ra)
    thu_muc_mo_hinh = THU_MUC / "models" / args.mo_hinh
    if not (thu_muc_mo_hinh / "model.bin").is_file():
        print(f"Chưa có mô hình nhận dạng tại {thu_muc_mo_hinh}")
        return 1
    files = tim_file_dau_vao(args.files, args.lam_lai)
    if not files:
        print(f"Không có file ghi âm mới. Chép file ghi âm vào {THU_MUC} rồi chạy lại.")
        return 2

    nap_thu_vien_cuda()
    print(f"Sẽ xử lý {len(files)} file. Đang nạp mô hình nhận dạng giọng nói...", flush=True)
    model, thiet_bi = nap_mo_hinh(args.cpu, thu_muc_mo_hinh)
    print(f"Chạy bằng: {thiet_bi}")

    so_loi = 0
    for i, f in enumerate(files, 1):
        print(f"\n{'=' * 70}\n({i}/{len(files)}) {f.name}\n{'=' * 70}", flush=True)
        try:
            model, thiet_bi = xu_ly_mot_file(f, model, thiet_bi, args.goi_y, thu_muc_mo_hinh)
        except KeyboardInterrupt:
            print("\nĐã dừng theo yêu cầu.")
            return 1
        except Exception as loi:
            so_loi += 1
            print(f"\nLỖI khi xử lý {f.name}: {loi}")
    return 1 if so_loi else 0


if __name__ == "__main__":
    sys.exit(main())
